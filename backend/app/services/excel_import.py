"""Ingesta del itinerario diario (equivalente a la hoja 'INFO' del Excel).

Reglas duras para no arrastrar los problemas del workbook original:
  - Se abre con openpyxl(data_only=True): solo se leen valores YA calculados
    por Excel, nunca se reevalúan fórmulas ni referencias externas
    ('[1]Hoja'!celda), que en el archivo original quedan rotas apenas se
    mueve o renombra el libro.
  - Se valida la presencia de encabezados esperados antes de aceptar el
    archivo; si falta una columna crítica se rechaza con un mensaje claro
    en vez de importar datos a medias.
  - Cada fila se valida por separado: una fila con hora ilegible o sin call
    sign se reporta como error pero no aborta el resto del archivo.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time

import openpyxl

from app.services.ctot_calc import parse_hhmm
from app.services.texto import normalizar

REQUIRED_HEADERS_ARR = ["CALL SIGN", "HORA  ARR UTC", "ORIGEN", "TIPO DE AERONAVE", "TIPO DE SERVICIO"]
REQUIRED_HEADERS_DEP = ["CALL SIGN", "HORA  DEP UTC", "DESTINO", "TIPO DE AERONAVE", "TIPO DE SERVICIO"]


@dataclass
class ItineraryRow:
    call_sign: str
    direction: str  # "ARR" | "DEP"
    hora_utc: str | None  # HHMM
    aerodromo: str | None
    tipo_aeronave: str | None
    tipo_servicio: str | None
    flight_date: date | None = None  # solo lo trae el formato 'temporada' (una fecha por fila)
    asientos: int | None = None
    # Numero de fila del archivo de origen. Sirve para volver del dato cargado
    # al renglon que lo produjo: hasta ahora solo los errores decian de que
    # fila venian, y las filas aceptadas quedaban sin procedencia.
    fila_origen: int | None = None


@dataclass
class HojaItinerario:
    """Una hoja del libro que tiene pinta de itinerario de temporada.

    El archivo que manda la DGAC es un libro de trabajo con varias hojas
    parciales (la convertida de la temporada pasada, la cruda "para
    convertir", la que arma el FMP). Se describen todas para que el operador
    vea cual se leyo y pueda elegir otra."""

    nombre: str
    fila_encabezado: int
    filas: int
    fecha_min: date | None = None
    fecha_max: date | None = None
    # Proporcion de filas de muestra cuyo numero de vuelo tiene pinta de call
    # sign OACI (SKX5368). En la hoja cruda de la DGAC el mismo campo trae
    # solo el numero ("5368"), y esa hoja no sirve para cargar.
    calidad: float = 0.0
    tiene_direccion: bool = False


@dataclass
class ImportResult:
    rows: list[ItineraryRow] = field(default_factory=list)
    errores: list[tuple[int, str]] = field(default_factory=list)  # (fila, motivo)
    hoja: str | None = None  # hoja del libro que se leyo (formato temporada)
    hojas: list[HojaItinerario] = field(default_factory=list)

    @property
    def filas_aceptadas(self) -> int:
        return len(self.rows)

    @property
    def filas_rechazadas(self) -> int:
        return len(self.errores)


class InvalidItineraryFile(Exception):
    pass


def _hhmm_from_cell(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, (time, datetime)):
        return f"{value.hour:02d}{value.minute:02d}"
    s = str(value).strip()
    if not s:
        return None
    # Segun como se haya armado la planilla, la misma columna viene como hora
    # de Excel o como texto ("00:05" / "00:05:00"). Se normaliza a HHMM, que es
    # lo que entiende parse_hhmm: sin esto, la hoja del itinerario de temporada
    # que usa el FMP se rechazaba fila por fila con "Hora ilegible".
    if ":" in s:
        partes = s.split(":")
        if len(partes) >= 2 and partes[0].strip().isdigit() and partes[1].strip().isdigit():
            return f"{int(partes[0]):02d}{int(partes[1]):02d}"
    return s


# Cuantas filas se miran, en cada hoja, buscando el encabezado. Los archivos
# de la DGAC no lo traen siempre en la fila 1: el libro de temporada abre con
# una fila de titulo ("ACTUALIZADO dd/mm/aaaa") y los recortes que arma el FMP
# quedan pegados mas abajo y corridos a la derecha.
MAX_FILAS_ENCABEZADO = 15

ALIAS_DIRECCION = ("SALIDA (D) LLEGADA (A)", "D / A", "D/A", "SALIDA / LLEGADA")


def _find_col(headers: list[str], *nombres: str, start: int = 0) -> int | None:
    """Indice de la primera columna (desde `start`) que coincide con alguno de
    los alias, o None.

    Cuidado con el patron anterior `find_col(a) or find_col(b)`: la columna 0
    es falsy, asi que un archivo con el dato en la primera columna --el
    "D / A" de la hoja que usa el FMP-- se leia como columna ausente."""
    for nombre in nombres:
        target = normalizar(nombre)
        for idx in range(start, len(headers)):
            if headers[idx] == target:
                return idx
    return None


def _buscar_encabezado(wb, puntaje, preferidas: tuple[str, ...] = ()):
    """Busca en que hoja y en que fila esta el encabezado del itinerario.

    Antes se daba por sentado que era la fila 1 (o la 2) de la primera hoja, y
    por eso se rechazaban archivos validos: el libro de la DGAC trae siete
    hojas y la buena no es la primera.

    `puntaje(headers)` devuelve 0 si esa fila no es un encabezado, 1 si le
    falta alguna columna util y 2 si esta completo; se queda con el mejor, asi
    que entre la hoja de trabajo a medio armar y la definitiva gana la
    definitiva. Devuelve (hoja, nro_de_fila, headers) o None."""
    orden = [n for n in preferidas if n in wb.sheetnames]
    orden += [n for n in wb.sheetnames if n not in orden]

    mejor = None  # (puntaje, hoja, fila, headers)
    for name in orden:
        for i, row in enumerate(wb[name].iter_rows(values_only=True), 1):
            if i > MAX_FILAS_ENCABEZADO:
                break
            headers = [normalizar(v) for v in row]
            valor = puntaje(headers)
            if valor and (mejor is None or valor > mejor[0]):
                mejor = (valor, name, i, headers)
                if valor == 2:
                    break
        if mejor is not None and mejor[0] == 2:
            break

    return None if mejor is None else (mejor[1], mejor[2], mejor[3])


def _puntaje_legacy(headers: list[str]) -> int:
    if _find_col(headers, "CALL SIGN") is None:
        return 0
    return 2 if _find_col(headers, "HORA ARR UTC") is not None else 1


def _puntaje_season(headers: list[str]) -> int:
    if _find_col(headers, "NUMERO DE VUELO") is None or _find_col(headers, "FECHA UTC") is None:
        return 0
    return 2 if _find_col(headers, *ALIAS_DIRECCION) is not None else 1


def _saltar_hasta(ws, fila: int):
    """Iterador de filas posicionado justo despues del encabezado."""
    rows_iter = ws.iter_rows(values_only=True)
    for _ in range(fila):
        next(rows_iter, None)
    return rows_iter


def parse_itinerary_excel(content: bytes) -> ImportResult:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # archivo corrupto / no es xlsx
        raise InvalidItineraryFile(f"No se pudo abrir el archivo: {exc}") from exc

    hallazgo = _buscar_encabezado(wb, _puntaje_legacy, preferidas=("INFO",))
    if hallazgo is None:
        raise InvalidItineraryFile(
            "No se encontraron las columnas de arribos (CALL SIGN / HORA ARR UTC) "
            "en ninguna hoja del archivo. Revisa que tenga el mismo formato que la hoja INFO."
        )
    sheet_name, fila_encabezado, headers = hallazgo
    ws = wb[sheet_name]

    col_arr_call = _find_col(headers, "CALL SIGN")
    col_arr_hora = _find_col(headers, "HORA ARR UTC")
    col_arr_aer = _find_col(headers, "ORIGEN")
    col_arr_tipo = _find_col(headers, "TIPO DE AERONAVE")
    col_arr_serv = _find_col(headers, "TIPO DE SERVICIO")

    if col_arr_call is None or col_arr_hora is None:
        raise InvalidItineraryFile(
            "No se encontraron las columnas de arribos (CALL SIGN / HORA ARR UTC) "
            "en la hoja de itinerario. Revisa que el archivo tenga el mismo formato que la hoja INFO."
        )

    col_dep_call = _find_col(headers, "CALL SIGN", start=col_arr_call + 1)
    col_dep_hora = _find_col(headers, "HORA DEP UTC", start=col_arr_call + 1)
    col_dep_aer = _find_col(headers, "DESTINO", start=col_arr_call + 1)
    col_dep_tipo = _find_col(headers, "TIPO DE AERONAVE", start=col_arr_call + 1)
    col_dep_serv = _find_col(headers, "TIPO DE SERVICIO", start=col_arr_call + 1)

    referencia = _referencia_externa(
        content, sheet_name, fila_encabezado, [col_arr_call, col_arr_hora]
    )
    if referencia is not None:
        raise _error_formulas_externas(referencia)

    rows_iter = _saltar_hasta(ws, fila_encabezado)
    result = ImportResult()
    row_num = fila_encabezado  # el numero que se reporta es el de la fila real del Excel

    for raw_row in rows_iter:
        row_num += 1

        def cell(idx: int | None):
            if idx is None or idx >= len(raw_row):
                return None
            return raw_row[idx]

        call_sign = cell(col_arr_call)
        if call_sign not in (None, ""):
            hora = _hhmm_from_cell(cell(col_arr_hora))
            if hora is not None and parse_hhmm(hora) is None:
                result.errores.append((row_num, f"Hora de arribo ilegible: {hora!r}"))
            else:
                result.rows.append(
                    ItineraryRow(
                        call_sign=str(call_sign).strip(),
                        direction="ARR",
                        hora_utc=hora,
                        aerodromo=str(cell(col_arr_aer) or "").strip() or None,
                        tipo_aeronave=str(cell(col_arr_tipo) or "").strip() or None,
                        tipo_servicio=str(cell(col_arr_serv) or "").strip() or None,
                        fila_origen=row_num,
                    )
                )

        if col_dep_call is not None:
            call_sign_dep = cell(col_dep_call)
            if call_sign_dep not in (None, ""):
                hora = _hhmm_from_cell(cell(col_dep_hora))
                if hora is not None and parse_hhmm(hora) is None:
                    result.errores.append((row_num, f"Hora de despegue ilegible: {hora!r}"))
                else:
                    result.rows.append(
                        ItineraryRow(
                            call_sign=str(call_sign_dep).strip(),
                            direction="DEP",
                            hora_utc=hora,
                            aerodromo=str(cell(col_dep_aer) or "").strip() or None,
                            tipo_aeronave=str(cell(col_dep_tipo) or "").strip() or None,
                            tipo_servicio=str(cell(col_dep_serv) or "").strip() or None,
                            fila_origen=row_num,
                        )
                    )

    return result


def detect_itinerary_format(content: bytes) -> str:
    """'season' = itinerario completo de temporada (una fecha por fila, IATA).
    'legacy' = carga diaria estilo hoja INFO (dos bloques ARR/DEP, OACI).

    Se busca el encabezado en todas las hojas, y no solo en la fila 1 de la
    primera: el libro de la DGAC abre con una fila de titulo y con hojas
    auxiliares, asi que un itinerario de temporada perfectamente valido se
    tomaba como 'legacy' y terminaba rechazado."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        return "legacy"
    return "season" if _buscar_encabezado(wb, _puntaje_season) else "legacy"


SEASON_DIRECTION_MAP = {"A": "ARR", "D": "DEP"}


MUESTRA_FORMULAS = 30  # filas que se miran buscando formulas a otro libro


def _es_formula_externa(texto: str) -> bool:
    """True para una formula que apunta a OTRO libro: '[1]Hoja'!celda.

    Se distingue de las referencias estructuradas de tabla (Tabla1[Columna])
    porque el corchete de un enlace externo encierra un numero."""
    i = texto.find("[")
    while i != -1:
        j = texto.find("]", i)
        if j > i + 1 and texto[i + 1 : j].isdigit():
            return True
        i = texto.find("[", i + 1)
    return False


def _tiene_enlaces_externos(content: bytes) -> bool:
    """Si el .xlsx no declara enlaces externos, no hace falta mirar formulas."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            return any(n.startswith("xl/externalLinks/") for n in z.namelist())
    except Exception:
        return False


def _referencia_externa(
    content: bytes, nombre_hoja: str, fila_encabezado: int, columnas: list
) -> str | None:
    """La primera formula a otro libro que aparezca en las columnas que
    importan (vuelo, fecha, hora), o None.

    Un archivo asi NO trae datos: lo que se ve en pantalla lo recalcula Excel
    al abrirlo, pero lo que queda guardado --lo unico que puede leer el
    sistema-- son los ultimos valores cacheados. Paso de verdad: un recorte
    "del 9 al 11 de setiembre" que por dentro guardaba el 1 al 4 de julio."""
    if not _tiene_enlaces_externos(content):
        return None
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=False, read_only=True)
    except Exception:
        return None
    if nombre_hoja not in wb.sheetnames:
        return None

    utiles = [c for c in columnas if c is not None]
    for n, raw in enumerate(_saltar_hasta(wb[nombre_hoja], fila_encabezado)):
        if n >= MUESTRA_FORMULAS:
            break
        for idx in utiles:
            if idx >= len(raw):
                continue
            valor = raw[idx]
            if isinstance(valor, str) and valor.startswith("=") and _es_formula_externa(valor):
                return valor
    return None


def _error_formulas_externas(referencia: str) -> InvalidItineraryFile:
    return InvalidItineraryFile(
        "El archivo no trae datos propios: sus celdas son fórmulas que apuntan a "
        f"otro libro ({referencia.lstrip('=')[:60]}). Excel muestra en pantalla lo "
        "que recalcula al abrirlo, pero lo que queda guardado --y lo único que "
        "puede leer el sistema-- son los últimos valores guardados, que pueden "
        "ser de otra fecha. Copiá la hoja y pegala como valores (Pegado especial "
        "→ Valores) y guardá, o subí directamente el libro original de la DGAC."
    )


MUESTRA_CALIDAD = 300  # filas que se miran para juzgar si la hoja sirve


def _pinta_de_call_sign(valor: object) -> bool:
    """"SKX5368" si; "5368" no. Sirve para distinguir la hoja ya convertida
    de la hoja cruda que la DGAC deja en el mismo libro."""
    texto = str(valor or "").strip().upper()
    return len(texto) >= 4 and texto[:3].isalpha() and any(c.isdigit() for c in texto[3:])


def _describir_hoja(ws, nombre: str, fila: int, headers: list[str]) -> HojaItinerario:
    """Recorre la hoja una vez para saber cuantos vuelos trae, que fechas
    cubre y si los numeros de vuelo tienen pinta de call sign."""
    col_vuelo = _find_col(headers, "NUMERO DE VUELO")
    col_fecha = _find_col(headers, "FECHA UTC")
    filas = muestra = validos = 0
    fecha_min = fecha_max = None

    for raw in _saltar_hasta(ws, fila):
        if col_vuelo is None or col_vuelo >= len(raw) or raw[col_vuelo] in (None, ""):
            continue
        filas += 1
        if muestra < MUESTRA_CALIDAD:
            muestra += 1
            validos += 1 if _pinta_de_call_sign(raw[col_vuelo]) else 0
        valor = raw[col_fecha] if col_fecha is not None and col_fecha < len(raw) else None
        fecha = valor.date() if isinstance(valor, datetime) else valor
        if isinstance(fecha, date):
            fecha_min = fecha if fecha_min is None or fecha < fecha_min else fecha_min
            fecha_max = fecha if fecha_max is None or fecha > fecha_max else fecha_max

    return HojaItinerario(
        nombre=nombre,
        fila_encabezado=fila,
        filas=filas,
        fecha_min=fecha_min,
        fecha_max=fecha_max,
        calidad=(validos / muestra) if muestra else 0.0,
        tiene_direccion=_find_col(headers, *ALIAS_DIRECCION) is not None,
    )


def _encabezado_season(ws) -> tuple[int, list[str]] | None:
    """(fila, encabezados) de la hoja, o None si no parece un itinerario."""
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i > MAX_FILAS_ENCABEZADO:
            break
        headers = [normalizar(v) for v in row]
        if _puntaje_season(headers):
            return i, headers
    return None


def _hojas_season(wb) -> list[tuple[HojaItinerario, list[str]]]:
    """Todas las hojas del libro que parecen un itinerario de temporada, ya
    descriptas (filas, fechas, calidad). Recorrerlas enteras cuesta unos
    segundos en el libro de la DGAC, asi que solo se hace cuando hay que
    elegir: si el operador ya eligio la hoja, se va directo a esa."""
    encontradas = []
    for nombre in wb.sheetnames:
        ws = wb[nombre]
        hallado = _encabezado_season(ws)
        if hallado is not None:
            fila, headers = hallado
            encontradas.append((_describir_hoja(ws, nombre, fila, headers), headers))
    return encontradas


def _elegir_hoja(candidatas: list[HojaItinerario], min_date: date | None) -> HojaItinerario | None:
    """La mejor hoja para la fecha pedida: primero las que llegan hasta esa
    fecha, y entre esas la que tiene los datos mas sanos.

    Con el libro real de la DGAC esto descarta la hoja de la temporada
    anterior (termina en junio) y la hoja cruda sin convertir (numeros de
    vuelo sin aerolinea), que es justo la que se elegia sola al tomar
    siempre la primera hoja del libro."""
    utiles = [h for h in candidatas if h.filas]
    if min_date is not None:
        utiles = [h for h in utiles if h.fecha_max is not None and h.fecha_max >= min_date]
    if not utiles:
        return None
    # Se prefiere una hoja con columna de direccion, pero si ninguna la tiene
    # igual se devuelve la mejor: asi el rechazo explica que falta esa columna,
    # que es el problema real, en vez de "ninguna hoja cubre la fecha".
    con_direccion = [h for h in utiles if h.tiene_direccion]
    return sorted(
        con_direccion or utiles, key=lambda h: (round(h.calidad, 2), h.filas), reverse=True
    )[0]


def parse_itinerary_season_excel(
    content: bytes,
    iata_to_icao: dict[str, str],
    min_date: date | None = None,
    hoja: str | None = None,
    estacion: str = "SPJC",
) -> ImportResult:
    """Formato 'itinerario completo' (temporada, ej. hoja 'W25'): una fila
    por vuelo por día, con fecha propia (columna 'FECHA UTC') y aeródromo
    en código IATA (columna 'ORIGEN / DESTINO') que se convierte a OACI
    usando `iata_to_icao`. Si `min_date` se especifica, se descartan (sin
    contar como error) las filas anteriores a esa fecha -- así se puede
    subir el archivo completo de la temporada y solo tomar "de tal fecha
    en adelante"."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise InvalidItineraryFile(f"No se pudo abrir el archivo: {exc}") from exc

    candidatas: list[tuple[HojaItinerario, list[str]]] = []

    if hoja is not None:
        # Camino corto: la hoja ya la eligio el operador en la vista previa.
        hallado = _encabezado_season(wb[hoja]) if hoja in wb.sheetnames else None
        if hallado is None:
            raise InvalidItineraryFile(
                f"La hoja {hoja!r} no esta en el archivo o no tiene formato de "
                "itinerario. Hojas del archivo: "
                + ", ".join(repr(n) for n in wb.sheetnames)
            )
        fila_encabezado, headers = hallado
        nombre_hoja = hoja
    else:
        candidatas = _hojas_season(wb)
        if not candidatas:
            raise InvalidItineraryFile(
                "No se encontraron las columnas del itinerario de temporada "
                "(NUMERO DE VUELO / FECHA UTC / ORIGEN / DESTINO) en ninguna hoja "
                "del archivo."
            )
        elegida = _elegir_hoja([h for h, _ in candidatas], min_date)
        if elegida is None:
            # Antes de culpar a las fechas: si la hoja son formulas a otro
            # libro, las fechas que trae guardadas no son las que se ven en
            # pantalla, y ese es el problema de verdad.
            for h, hs in candidatas:
                ref = _referencia_externa(
                    content,
                    h.nombre,
                    h.fila_encabezado,
                    [_find_col(hs, "NUMERO DE VUELO"), _find_col(hs, "FECHA UTC")],
                )
                if ref is not None:
                    raise _error_formulas_externas(ref)
            detalle = "; ".join(
                f"{h.nombre!r} va del {h.fecha_min} al {h.fecha_max}"
                for h, _ in candidatas
                if h.fecha_min is not None
            )
            raise InvalidItineraryFile(
                f"Ninguna hoja del archivo tiene vuelos desde el {min_date}. "
                + (f"Lo que trae el archivo: {detalle}." if detalle else "")
            )
        nombre_hoja = elegida.nombre
        fila_encabezado = elegida.fila_encabezado
        headers = dict((h.nombre, hs) for h, hs in candidatas)[nombre_hoja]

    ws = wb[nombre_hoja]
    rows_iter = _saltar_hasta(ws, fila_encabezado)

    col_vuelo = _find_col(headers, "NUMERO DE VUELO")
    col_dir = _find_col(headers, *ALIAS_DIRECCION)
    col_asientos = _find_col(headers, "ASIENTOS")
    col_tipo_ac = _find_col(headers, "TIPO DE AERONAVE")
    col_aerodromo = _find_col(headers, "ORIGEN / DESTINO")
    col_fecha = _find_col(headers, "FECHA UTC")
    col_hora = _columna_hora_estacion(headers, estacion, iata_to_icao)
    col_serv = _find_col(headers, "TIPO DE SERVICIO")

    if col_vuelo is None or col_fecha is None or col_aerodromo is None:
        raise InvalidItineraryFile(
            "No se encontraron las columnas esperadas del itinerario de temporada "
            "(NÚMERO DE VUELO / FECHA UTC / ORIGEN / DESTINO)."
        )

    referencia = _referencia_externa(
        content, nombre_hoja, fila_encabezado, [col_vuelo, col_fecha, col_hora]
    )
    if referencia is not None:
        raise _error_formulas_externas(referencia)

    # Sin columna de direccion no hay forma de saber si la fila es arribo o
    # despegue: se corta aca con un mensaje claro, en vez de reportar
    # "direccion invalida" en cada una de las miles de filas del archivo.
    if col_dir is None:
        raise InvalidItineraryFile(
            'Al archivo le falta la columna de dirección ("D / A" o "SALIDA (D) / '
            'LLEGADA (A)"), así que no se puede saber si cada vuelo es arribo o '
            "despegue. Suele pasar cuando se copian solo algunas columnas del "
            "itinerario en una hoja nueva: cargue el archivo completo de la DGAC."
        )

    result = ImportResult(hoja=nombre_hoja, hojas=[h for h, _ in candidatas])
    row_num = fila_encabezado  # se reporta el numero de fila real del Excel

    for raw_row in rows_iter:
        row_num += 1

        def cell(idx: int | None):
            if idx is None or idx >= len(raw_row):
                return None
            return raw_row[idx]

        call_sign = cell(col_vuelo)
        if call_sign in (None, ""):
            continue

        fecha_val = cell(col_fecha)
        flight_date = fecha_val.date() if isinstance(fecha_val, datetime) else fecha_val
        if not isinstance(flight_date, date):
            result.errores.append((row_num, f"Fecha ilegible: {fecha_val!r}"))
            continue
        if min_date is not None and flight_date < min_date:
            continue

        direccion_raw = str(cell(col_dir) or "").strip().upper()
        direction = SEASON_DIRECTION_MAP.get(direccion_raw)
        if direction is None:
            result.errores.append((row_num, f"Dirección inválida (debe ser A/D): {direccion_raw!r}"))
            continue

        hora = _hhmm_from_cell(cell(col_hora))
        if hora is None or parse_hhmm(hora) is None:
            result.errores.append((row_num, f"Hora ilegible: {hora!r}"))
            continue

        iata = str(cell(col_aerodromo) or "").strip().upper()
        icao = iata_to_icao.get(iata) or iata or None

        asientos_val = cell(col_asientos)
        try:
            asientos = int(asientos_val) if asientos_val not in (None, "") else None
        except (TypeError, ValueError):
            asientos = None

        result.rows.append(
            ItineraryRow(
                call_sign=str(call_sign).strip(),
                direction=direction,
                hora_utc=hora,
                aerodromo=icao,
                tipo_aeronave=str(cell(col_tipo_ac) or "").strip() or None,
                tipo_servicio=str(cell(col_serv) or "").strip() or None,
                flight_date=flight_date,
                asientos=asientos,
                fila_origen=row_num,
            )
        )

    return result


def _columna_hora_estacion(
    headers: list[str], estacion: str, iata_to_icao: dict[str, str]
) -> int:
    """La hora aprobada pertenece al aeropuerto que recibe la carga.

    Un encabezado genérico debe declarar UTC. Los encabezados que nombran
    aeropuerto se contrastan con el catálogo, incluyendo el formato antiguo
    «HORA APROBADA LIM» que ya utiliza el sistema.
    """
    estacion = estacion.strip().upper()
    codigos = {estacion} | {
        iata.strip().upper() for iata, oaci in iata_to_icao.items()
        if oaci and oaci.strip().upper() == estacion
    }
    columnas: list[int] = []
    for indice, encabezado in enumerate(headers):
        if encabezado in {"HORA UTC", "HORA APROBADA UTC"}:
            columnas.append(indice)
            continue
        explicito = re.fullmatch(r"HORA APROBADA ([A-Z]{3,4})(?: UTC)?", encabezado)
        if explicito:
            codigo = explicito.group(1)
            if codigo not in codigos:
                raise InvalidItineraryFile(
                    f"La columna {encabezado!r} corresponde a {codigo}, "
                    f"pero la carga está seleccionada para {estacion}. "
                    "Seleccione el aeropuerto correcto o revise el archivo."
                )
            columnas.append(indice)
    if not columnas:
        raise InvalidItineraryFile(
            f"Falta la columna de hora aprobada de {estacion}. Use "
            "HORA APROBADA <código IATA/OACI> UTC o HORA APROBADA UTC."
        )
    if len(columnas) != 1:
        raise InvalidItineraryFile(
            "Hay varias columnas de hora aprobada; deje una sola para evitar "
            "importar un horario ambiguo."
        )
    return columnas[0]


def parse_itinerary_csv(content: bytes) -> ImportResult:
    """Esquema plano alternativo: call_sign,direction,hora_utc,aerodromo,tipo_aeronave,tipo_servicio"""
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    required = {"call_sign", "direction", "hora_utc"}
    if reader.fieldnames is None or not required.issubset({f.strip().lower() for f in reader.fieldnames}):
        raise InvalidItineraryFile(
            f"El CSV debe tener al menos las columnas: {', '.join(sorted(required))}"
        )

    result = ImportResult()
    for row_num, raw in enumerate(reader, start=2):
        row = {k.strip().lower(): v for k, v in raw.items()}
        call_sign = (row.get("call_sign") or "").strip()
        direction = (row.get("direction") or "").strip().upper()
        hora = (row.get("hora_utc") or "").strip() or None

        if not call_sign:
            result.errores.append((row_num, "Falta call sign"))
            continue
        if direction not in ("ARR", "DEP"):
            result.errores.append((row_num, f"direction inválido (debe ser ARR/DEP): {direction!r}"))
            continue
        if hora is not None and parse_hhmm(hora) is None:
            result.errores.append((row_num, f"Hora ilegible: {hora!r}"))
            continue

        result.rows.append(
            ItineraryRow(
                call_sign=call_sign,
                direction=direction,
                hora_utc=hora,
                aerodromo=(row.get("aerodromo") or "").strip() or None,
                tipo_aeronave=(row.get("tipo_aeronave") or "").strip() or None,
                tipo_servicio=(row.get("tipo_servicio") or "").strip() or None,
                fila_origen=row_num,
            )
        )
    return result
