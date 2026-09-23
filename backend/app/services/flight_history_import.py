"""Ingesta de vuelos históricos: una grilla FormatoFMP ya cerrada, de meses o
años anteriores.

Entran dos formatos, y el punto de entrada es `parse_flight_history_file`:

  - El **CONSOLIDADO del día** (.xlsx), que es lo que la FMP archiva de
    verdad: un libro por día con una hoja por sector. Es el formato que el
    operador tiene a mano.
  - El **CSV** exportado de una hoja, que existe solo cuando alguien lo
    genera aparte.

Ese CSV es el export crudo de un Excel con columnas duplicadas (un bloque de
verificación por fórmulas al lado del real) y encabezados que no siempre usan
el mismo espaciado entre archivos ("ETD 1" vs "ETD2"). Por eso las columnas
se ubican por nombre normalizado (mayúsculas, sin tildes, espacios
colapsados) y con una segunda pasada ignorando espacios del todo, en vez de
por índice fijo -- así no se rompe si un archivo de otro mes corre las
columnas o cambia el espaciado.

flight_date y sector NO se leen del archivo (el CSV no trae una fecha
confiable por fila -- ver notas de la conversación): los indica quien sube
el archivo, uno por día.
"""
from __future__ import annotations

import csv
import datetime
import io
from dataclasses import dataclass, field

import openpyxl

from app.services.texto import normalizar

REQUIRED_HEADERS = ["VUELO", "HORA"]

# Límites de largo de columna en el modelo Flight (backend/app/models.py) --
# se trunca en vez de fallar, porque son datos históricos de solo lectura y
# no vale la pena rechazar la fila entera por un campo de texto libre largo.
_MAX_LEN = {
    "hora": 4, "d_ats": 16, "vuelo": 16, "adep": 8, "dep": 4,
    "eta_aircon": 4, "eta_aircon_calculado": 4, "slot_arr_dgac": 16,
    "etd1": 4, "ctot1": 4, "sec1": 16, "h_rev1": 4, "rev1": 16,
    "etd2": 4, "ctot2": 4, "sec2": 16, "h_rev2": 4, "rev2": 16,
    "etd3": 4, "ctot3": 4, "sec3": 16, "h_rev3": 4, "rev3": 16,
    "etd4": 4, "ctot4": 4, "sec4": 16,
    "observaciones": 500,
}

# Placeholders del Excel viejo que en el sistema equivalen a "vacío".
_EMPTY_PLACEHOLDERS = {"NO H. ITIN", "N/A", "NA", "-", "--"}


@dataclass
class FlightHistoryRow:
    vuelo: str
    hora: str | None = None
    d_ats: str | None = None
    adep: str | None = None
    dep: str | None = None
    eta_aircon: str | None = None
    eta_aircon_calculado: str | None = None
    slot_arr_dgac: str | None = None
    etd1: str | None = None
    ctot1: str | None = None
    sec1: str | None = None
    h_rev1: str | None = None
    rev1: str | None = None
    etd2: str | None = None
    ctot2: str | None = None
    sec2: str | None = None
    h_rev2: str | None = None
    rev2: str | None = None
    etd3: str | None = None
    ctot3: str | None = None
    sec3: str | None = None
    h_rev3: str | None = None
    rev3: str | None = None
    etd4: str | None = None
    ctot4: str | None = None
    sec4: str | None = None
    dla_minutos: float | None = None
    observaciones: str | None = None
    cancelado: bool = False
    # Numero de fila del archivo de origen, para poder volver del vuelo
    # cargado al renglon que lo produjo.
    fila_origen: int | None = None


def campos_del_vuelo_historico(row: FlightHistoryRow) -> dict[str, str | float | bool | None]:
    """Valores operacionales del archivo, sin normalizarlos ni recalcularlos.

    El importador aporta por separado fecha, sector, autor y procedencia.
    No depende del ORM: también se usa desde el script de carga histórica.
    """
    return {
        "hora": row.hora, "d_ats": row.d_ats, "vuelo": row.vuelo,
        "adep": row.adep, "dep": row.dep,
        "eta_aircon": row.eta_aircon, "eta_aircon_calculado": row.eta_aircon_calculado,
        "slot_arr_dgac": row.slot_arr_dgac,
        "etd1": row.etd1, "ctot1": row.ctot1, "sec1": row.sec1,
        "h_rev1": row.h_rev1, "rev1": row.rev1,
        "etd2": row.etd2, "ctot2": row.ctot2, "sec2": row.sec2,
        "h_rev2": row.h_rev2, "rev2": row.rev2,
        "etd3": row.etd3, "ctot3": row.ctot3, "sec3": row.sec3,
        "h_rev3": row.h_rev3, "rev3": row.rev3,
        "etd4": row.etd4, "ctot4": row.ctot4, "sec4": row.sec4,
        "dla_minutos": row.dla_minutos,
        "observaciones": row.observaciones, "cancelado": row.cancelado,
    }


@dataclass
class FlightHistoryImportResult:
    rows: list[FlightHistoryRow] = field(default_factory=list)
    errores: list[tuple[int, str]] = field(default_factory=list)

    @property
    def filas_aceptadas(self) -> int:
        return len(self.rows)

    @property
    def filas_rechazadas(self) -> int:
        return len(self.errores)


class InvalidFlightHistoryFile(Exception):
    pass


def _norm_nospace(s: object) -> str:
    return normalizar(s).replace(" ", "")


def _hhmm(value: str | None) -> str | None:
    """Normaliza 'HH:MM' -> 'HHMM'; deja pasar el resto tal cual (placeholders
    de texto como 'NO H. ITIN' se limpian aparte, ver _clean_slot)."""
    if value is None:
        return None
    v = value.strip()
    if not v:
        return None
    if ":" in v:
        hh, _, mm = v.partition(":")
        v = f"{hh.strip().zfill(2)}{mm.strip().zfill(2)}"
    return v


def _clean_text(value: str | None) -> str | None:
    if value is None:
        return None
    v = value.strip()
    return v or None


def _clean_slot(value: str | None) -> str | None:
    v = _clean_text(value)
    if v is None:
        return None
    return None if normalizar(v) in _EMPTY_PLACEHOLDERS else _hhmm(v)


def _truncate(field_name: str, value: str | None) -> str | None:
    if value is None:
        return None
    limit = _MAX_LEN.get(field_name)
    return value[:limit] if limit else value


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    # El mensaje decía solo "codificación desconocida", y el caso real que lo
    # disparaba no era una codificación rara: era un archivo que no es texto.
    # Un .xlsx es un ZIP, y sus bytes chocan casi siempre con alguno de los
    # cinco que cp1252 no define -- así que el operador que soltaba el
    # CONSOLIDADO del día recibía un error sobre codificaciones y no tenía
    # forma de saber qué le estaban pidiendo.
    raise InvalidFlightHistoryFile(
        "No se pudo leer el archivo: no parece un CSV ni un libro de Excel. "
        "Se esperan el CONSOLIDADO del día (.xlsx) o el CSV exportado de la "
        "grilla."
    )


def parse_flight_history_csv(content: bytes) -> FlightHistoryImportResult:
    text = _decode(content)
    all_rows = list(csv.reader(io.StringIO(text), delimiter=";"))

    header_idx = None
    for idx, raw_row in enumerate(all_rows[:20]):
        normed = {normalizar(c) for c in raw_row}
        if "VUELO" in normed and "HORA" in normed:
            header_idx = idx
            break
    if header_idx is None:
        raise InvalidFlightHistoryFile(
            "No se encontró la fila de encabezados (se busca VUELO y HORA) "
            "entre las primeras filas del archivo."
        )

    headers = all_rows[header_idx]
    normed_headers = [normalizar(h) for h in headers]
    nospace_headers = [_norm_nospace(h) for h in headers]

    def find_col(*candidates: str) -> int | None:
        for cand in candidates:
            target = normalizar(cand)
            if target in normed_headers:
                return normed_headers.index(target)
        for cand in candidates:
            target = _norm_nospace(cand)
            if target in nospace_headers:
                return nospace_headers.index(target)
        return None

    columns = {
        "vuelo": find_col("VUELO"),
        "hora": find_col("HORA"),
        "d_ats": find_col("D. ATS", "D.ATS"),
        "adep": find_col("ADEP"),
        "dep": find_col("DEP"),
        "eta_aircon": find_col("ETA AIRCON"),
        "eta_aircon_calculado": find_col("ETA AIRCON CALCULADO"),
        "slot_arr_dgac": find_col("SLOT ARR DGAC"),
        "etd1": find_col("ETD 1", "ETD1"),
        "ctot1": find_col("CTOT 1", "CTOT1"),
        "sec1": find_col("SEC1"),
        "h_rev1": find_col("H. REV1", "H.REV1"),
        "rev1": find_col("REV1"),
        "etd2": find_col("ETD 2", "ETD2"),
        "ctot2": find_col("CTOT 2", "CTOT2"),
        "sec2": find_col("SEC2"),
        "h_rev2": find_col("H. REV2", "H.REV2"),
        "rev2": find_col("REV2"),
        "etd3": find_col("ETD 3", "ETD3"),
        "ctot3": find_col("CTOT 3", "CTOT3"),
        "sec3": find_col("SEC3"),
        "h_rev3": find_col("H. REV3", "H.REV3"),
        "rev3": find_col("REV3"),
        "etd4": find_col("ETD 4", "ETD4"),
        "ctot4": find_col("CTOT 4", "CTOT4"),
        "sec4": find_col("SEC4"),
        "dla_minutos": find_col("DLA (MIN)", "DLA MIN"),
        "observaciones": find_col("OBSERVACIONES"),
    }

    if columns["vuelo"] is None or columns["hora"] is None:
        raise InvalidFlightHistoryFile(
            "No se encontraron las columnas VUELO / HORA en el encabezado del archivo."
        )

    # Las columnas ETD/CTOT/etc. se repiten más a la derecha en un bloque de
    # verificación por fórmulas (mismos nombres, valores TRUE/FALSE u horas
    # con otro formato) -- si find_col() encontró la primera aparición (la
    # real) esto ya no importa, pero se limita la búsqueda a antes de esa
    # segunda aparición de VUELO/HORA para mayor seguridad.
    second_vuelo = None
    for i in range(columns["vuelo"] + 1, len(headers)):
        if normed_headers[i] == "VUELO":
            second_vuelo = i
            break
    boundary = second_vuelo if second_vuelo is not None else len(headers)
    for key, idx in columns.items():
        if idx is not None and idx >= boundary:
            columns[key] = None

    result = FlightHistoryImportResult()
    data_rows = all_rows[header_idx + 1 :]

    def cell(raw_row: list[str], idx: int | None) -> str | None:
        if idx is None or idx >= len(raw_row):
            return None
        return raw_row[idx]

    for row_num, raw_row in enumerate(data_rows, start=header_idx + 2):
        vuelo = _clean_text(cell(raw_row, columns["vuelo"]))
        if not vuelo:
            continue  # fila vacía (de sobra al final del archivo): se ignora, no es error

        try:
            dla_raw = cell(raw_row, columns["dla_minutos"])
            dla_raw = dla_raw.strip() if dla_raw else ""
            dla_minutos = float(dla_raw.replace(",", ".")) if dla_raw else None
        except ValueError:
            result.errores.append((row_num, f"DLA (min) ilegible: {dla_raw!r}"))
            continue

        observaciones = _truncate("observaciones", _clean_text(cell(raw_row, columns["observaciones"])))
        cancelado = bool(observaciones) and "CANCELADO" in normalizar(observaciones)

        row = FlightHistoryRow(
            fila_origen=row_num,
            vuelo=_truncate("vuelo", vuelo),
            hora=_truncate("hora", _hhmm(cell(raw_row, columns["hora"]))),
            d_ats=_truncate("d_ats", _clean_text(cell(raw_row, columns["d_ats"]))),
            adep=_truncate("adep", _clean_text(cell(raw_row, columns["adep"]))),
            dep=_truncate("dep", _hhmm(cell(raw_row, columns["dep"]))),
            eta_aircon=_truncate("eta_aircon", _hhmm(cell(raw_row, columns["eta_aircon"]))),
            eta_aircon_calculado=_truncate("eta_aircon_calculado", _hhmm(cell(raw_row, columns["eta_aircon_calculado"]))),
            slot_arr_dgac=_truncate("slot_arr_dgac", _clean_slot(cell(raw_row, columns["slot_arr_dgac"]))),
            etd1=_truncate("etd1", _hhmm(cell(raw_row, columns["etd1"]))),
            ctot1=_truncate("ctot1", _hhmm(cell(raw_row, columns["ctot1"]))),
            sec1=_truncate("sec1", _clean_text(cell(raw_row, columns["sec1"]))),
            h_rev1=_truncate("h_rev1", _hhmm(cell(raw_row, columns["h_rev1"]))),
            rev1=_truncate("rev1", _clean_text(cell(raw_row, columns["rev1"]))),
            etd2=_truncate("etd2", _hhmm(cell(raw_row, columns["etd2"]))),
            ctot2=_truncate("ctot2", _hhmm(cell(raw_row, columns["ctot2"]))),
            sec2=_truncate("sec2", _clean_text(cell(raw_row, columns["sec2"]))),
            h_rev2=_truncate("h_rev2", _hhmm(cell(raw_row, columns["h_rev2"]))),
            rev2=_truncate("rev2", _clean_text(cell(raw_row, columns["rev2"]))),
            etd3=_truncate("etd3", _hhmm(cell(raw_row, columns["etd3"]))),
            ctot3=_truncate("ctot3", _hhmm(cell(raw_row, columns["ctot3"]))),
            sec3=_truncate("sec3", _clean_text(cell(raw_row, columns["sec3"]))),
            h_rev3=_truncate("h_rev3", _hhmm(cell(raw_row, columns["h_rev3"]))),
            rev3=_truncate("rev3", _clean_text(cell(raw_row, columns["rev3"]))),
            etd4=_truncate("etd4", _hhmm(cell(raw_row, columns["etd4"]))),
            ctot4=_truncate("ctot4", _hhmm(cell(raw_row, columns["ctot4"]))),
            sec4=_truncate("sec4", _clean_text(cell(raw_row, columns["sec4"]))),
            dla_minutos=dla_minutos,
            observaciones=observaciones,
            cancelado=cancelado,
        )
        result.rows.append(row)

    return result


# --------------------------------------------------------------------------- #
# El libro CONSOLIDADO del día
#
# Lo que la FMP archiva no es el CSV: es un Excel por día con una hoja por
# sector ('FormatoFMP SUR', 'FormatoFMP NOR') y el itinerario en 'INFO'. El
# CSV existe solo cuando alguien exporta una hoja a mano.
#
# Leerlo vivía únicamente en scripts/import_historico.py, así que la carga por
# pantalla no podía con el archivo que el operador tiene de verdad: soltaba el
# CONSOLIDADO y recibía un error de codificación. Vive acá para que el script
# y el endpoint usen el mismo lector, igual que ya pasó con `texto.normalizar`.
# --------------------------------------------------------------------------- #

#: Los primeros bytes de todo ZIP, y un .xlsx es un ZIP.
FIRMA_ZIP = b"PK\x03\x04"

#: Valores de error de fórmula: se tratan como celda vacía.
ERRORES_EXCEL = {"#REF!", "#VALUE!", "#N/A", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!"}

#: Prefijo de la hoja de cada sector dentro del libro.
HOJA_DE_SECTOR = {"SUR": "FormatoFMP SUR", "NOR": "FormatoFMP NOR"}


def es_libro_excel(content: bytes) -> bool:
    return content[:4] == FIRMA_ZIP


def celda_a_texto(valor: object) -> str:
    """Rinde una celda de la grilla como el texto que el parser sabe leer.

    Las horas del Excel vienen como `datetime.time` ('HH:MM:SS') y se pasan a
    'HH:MM' para que `_hhmm` las normalice a 'HHMM'. Los flotantes enteros
    (2305.0) se limpian y el ruido de coma flotante del DLA (2.00000..1) se
    redondea. Los errores de fórmula se tratan como vacío: una celda con #REF!
    no es un dato, y escribirla tal cual metería '#REF!' en la base.
    """
    if valor is None:
        return ""
    if isinstance(valor, str) and valor.strip().upper() in ERRORES_EXCEL:
        return ""
    if isinstance(valor, (datetime.datetime, datetime.time)):
        return valor.strftime("%H:%M")
    if isinstance(valor, float):
        valor = round(valor, 2)
        return str(int(valor)) if valor.is_integer() else str(valor)
    return str(valor)


def buscar_hoja(wb, *prefijos: str) -> str | None:
    """La primera hoja cuyo nombre normalizado empieza con alguno de los
    prefijos. Por prefijo y no por igualdad: hay libros con 'FormatoFMP NORTE'
    y con espacios de más, y son la misma hoja."""
    objetivos = [normalizar(p) for p in prefijos]
    for nombre in wb.sheetnames:
        normalizado = normalizar(nombre)
        if any(normalizado.startswith(objetivo) for objetivo in objetivos):
            return nombre
    return None


def hoja_a_csv(ws) -> bytes:
    """La hoja, volcada al mismo CSV que el parser ya sabe leer."""
    buf = io.StringIO()
    escritor = csv.writer(buf, delimiter=";")
    for fila in ws.iter_rows(values_only=True):
        escritor.writerow([celda_a_texto(celda) for celda in fila])
    return buf.getvalue().encode("utf-8-sig")


def parse_flight_history_excel(content: bytes, sector: str) -> FlightHistoryImportResult:
    """Extrae del libro la hoja del sector pedido y la parsea.

    El libro trae los dos sectores; cuál cargar lo dice quien sube el archivo,
    igual que la fecha. Tomar los dos de una sería otra operación: el endpoint
    carga un sector y un día.
    """
    prefijo = HOJA_DE_SECTOR.get(str(sector).upper())
    if prefijo is None:
        raise InvalidFlightHistoryFile(f"Sector desconocido: {sector}")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl levanta de todo ante un ZIP que no es xlsx
        raise InvalidFlightHistoryFile(
            "No se pudo abrir el libro de Excel. Si es un .xls antiguo, hay que "
            "volver a guardarlo como .xlsx."
        ) from exc

    try:
        hoja = buscar_hoja(wb, prefijo)
        if hoja is None:
            disponibles = ", ".join(wb.sheetnames) or "ninguna"
            raise InvalidFlightHistoryFile(
                f"El libro no tiene la hoja '{prefijo}' del sector {sector}. "
                f"Hojas encontradas: {disponibles}."
            )
        return parse_flight_history_csv(hoja_a_csv(wb[hoja]))
    finally:
        # read_only deja el archivo abierto hasta que se cierra explícitamente.
        wb.close()


def parse_flight_history_file(content: bytes, sector: str) -> FlightHistoryImportResult:
    """Punto de entrada único: acepta el CONSOLIDADO del día o el CSV.

    El formato se decide por el contenido y no por la extensión del nombre:
    un archivo renombrado a .csv sigue siendo un libro de Excel, y el operador
    no tiene por qué saber cuál de los dos le toca.
    """
    if es_libro_excel(content):
        return parse_flight_history_excel(content, sector)
    return parse_flight_history_csv(content)
