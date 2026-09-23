from datetime import date
from pathlib import Path

import openpyxl
import pytest

from app.airport_codes import SEED_AIRPORTS
from app.services.excel_import import (
    InvalidItineraryFile,
    detect_itinerary_format,
    parse_itinerary_csv,
    parse_itinerary_excel,
    parse_itinerary_season_excel,
)

# Ruta relativa al propio test: antes eran rutas absolutas de la maquina de
# desarrollo, asi que en cualquier otra computadora los `skipif` de abajo
# saltaban estos tests EN SILENCIO -- la suite pasaba en verde sin haber
# comprobado nada contra los archivos reales.
FIXTURES = Path(__file__).parent / "fixtures"
SAMPLE_XLSX = FIXTURES / "JUNIO 3 CONSOLIDADO(1).xlsx"
SEASON_XLSX = FIXTURES / "INTINERARIO.xlsx"
IATA_TO_ICAO = {code: info["icao"] for code, info in SEED_AIRPORTS.items() if info["icao"]}


@pytest.mark.skipif(not SAMPLE_XLSX.exists(), reason="archivo de muestra no disponible")
def test_parse_real_info_sheet():
    result = parse_itinerary_excel(SAMPLE_XLSX.read_bytes())
    assert result.filas_aceptadas > 0
    arrivals = [r for r in result.rows if r.direction == "ARR"]
    departures = [r for r in result.rows if r.direction == "DEP"]
    assert arrivals and departures
    first_arr = arrivals[0]
    assert first_arr.call_sign == "LPE2121"
    assert first_arr.hora_utc == "0010"
    assert first_arr.aerodromo == "SPZO"
    assert first_arr.tipo_aeronave == "A320"


def test_parse_csv_basic():
    csv_content = (
        "call_sign,direction,hora_utc,aerodromo,tipo_aeronave,tipo_servicio\n"
        "LPE123,ARR,0100,SPZO,A320,J\n"
        "LPE124,DEP,0200,SPQU,A320,J\n"
    ).encode("utf-8")
    result = parse_itinerary_csv(csv_content)
    assert result.filas_aceptadas == 2
    assert result.filas_rechazadas == 0


def test_parse_csv_missing_columns_rejected():
    csv_content = "foo,bar\n1,2\n".encode("utf-8")
    with pytest.raises(InvalidItineraryFile):
        parse_itinerary_csv(csv_content)


def test_parse_csv_bad_row_reported_not_fatal():
    csv_content = (
        "call_sign,direction,hora_utc\n"
        "LPE123,ARR,9999\n"
        "LPE124,DEP,0200\n"
    ).encode("utf-8")
    result = parse_itinerary_csv(csv_content)
    assert result.filas_aceptadas == 1
    assert result.filas_rechazadas == 1


def _build_season_xlsx(rows: list[tuple]) -> bytes:
    import io

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "NÚMERO DE VUELO", "SALIDA (D)\nLLEGADA (A)", "ASIENTOS", "TIPO DE AERONAVE",
        "ORIGEN / DESTINO", "VIA", "FECHA UTC", "HORA APROBADA LIM UTC", "TIPO DE SERVICIO",
    ])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_detect_format_season_vs_legacy():
    season_bytes = _build_season_xlsx([
        ("LPE2000", "D", 174, "A320", "TYL", "TYL", date(2026, 1, 1), None, "J"),
    ])
    assert detect_itinerary_format(season_bytes) == "season"
    assert detect_itinerary_format(SAMPLE_XLSX.read_bytes()) == "legacy"


def test_parse_season_format_converts_iata_and_keeps_per_row_date():
    from datetime import time as time_

    content = _build_season_xlsx([
        ("LPE2000", "D", 174, "A320", "TYL", "TYL", date(2026, 1, 1), time_(0, 0), "J"),
        ("JAP7221", "A", 186, "A320", "PIU", "PIU", date(2026, 1, 2), time_(0, 5), "J"),
    ])
    result = parse_itinerary_season_excel(content, IATA_TO_ICAO)
    assert result.filas_aceptadas == 2
    dep, arr = result.rows
    assert dep.call_sign == "LPE2000"
    assert dep.direction == "DEP"
    assert dep.aerodromo == "SPYL"  # TYL -> SPYL (Talara)
    assert dep.flight_date == date(2026, 1, 1)
    assert dep.asientos == 174
    assert arr.aerodromo == "SPUR"  # PIU -> SPUR (Piura)
    assert arr.flight_date == date(2026, 1, 2)


def test_parse_season_format_min_date_filters_past_rows():
    from datetime import time as time_

    content = _build_season_xlsx([
        ("LPE1", "D", 100, "A320", "CUZ", "CUZ", date(2026, 1, 1), time_(0, 0), "J"),
        ("LPE2", "D", 100, "A320", "CUZ", "CUZ", date(2026, 7, 4), time_(0, 0), "J"),
    ])
    result = parse_itinerary_season_excel(content, IATA_TO_ICAO, min_date=date(2026, 7, 4))
    assert result.filas_aceptadas == 1
    assert result.rows[0].call_sign == "LPE2"


@pytest.mark.skipif(not SEASON_XLSX.exists(), reason="itinerario completo no disponible")
def test_parse_real_season_itinerary():
    result = parse_itinerary_season_excel(
        SEASON_XLSX.read_bytes(), IATA_TO_ICAO, min_date=date(2026, 1, 1)
    )
    assert result.filas_aceptadas > 200_000
    sample = result.rows[0]
    assert sample.call_sign == "LPE2000"
    assert sample.direction == "DEP"
    assert sample.aerodromo == "SPYL"  # TYL -> SPYL
    assert sample.flight_date == date(2026, 1, 1)


# --- Casos reales del archivo que manda la DGAC (set-2026) -------------------
# El libro no viene como se suponia: el encabezado no siempre esta en la fila 1,
# el libro trae varias hojas parciales (una de la temporada anterior, otra sin
# convertir) y las horas a veces son texto "HH:MM" en vez de hora de Excel.


def _libro(hojas: dict, encabezado_en: dict | None = None) -> bytes:
    """Libro con varias hojas; `encabezado_en` corre el encabezado de una hoja
    a la fila indicada (y a la columna C) para imitar los recortes pegados."""
    import io

    encabezado_en = encabezado_en or {}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre, (headers, filas) in hojas.items():
        ws = wb.create_sheet(nombre)
        desde = encabezado_en.get(nombre)
        if desde:
            for _ in range(desde - 1):
                ws.append([])
            relleno = ["", ""]
        else:
            relleno = []
        ws.append(relleno + list(headers))
        for f in filas:
            ws.append(relleno + list(f))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HEADERS_TEMPORADA = [
    "NÚMERO DE VUELO", "SALIDA (D)" + chr(10) + "LLEGADA (A)", "ASIENTOS", "TIPO DE AERONAVE",
    "ORIGEN / DESTINO", "VIA", "FECHA UTC", "HORA APROBADA LIM UTC", "TIPO DE SERVICIO",
]
HEADERS_FMP = [
    "D / A", "NÚMERO DE VUELO", "TIPO DE AERONAVE", "ORIGEN / DESTINO",
    "FECHA UTC", "HORA APROBADA LIM UTC", "TIPO DE SERVICIO",
]


def test_encabezado_fuera_de_la_fila_1_igual_se_encuentra():
    """El recorte que arma el FMP queda pegado en C5, no en A1."""
    content = _libro(
        {"Hoja1": (HEADERS_TEMPORADA, [
            ("LPE2000", "D", 174, "A320", "TYL", "TYL", date(2026, 9, 9), "00:05", "J"),
        ])},
        encabezado_en={"Hoja1": 5},
    )
    assert detect_itinerary_format(content) == "season"
    result = parse_itinerary_season_excel(content, IATA_TO_ICAO)
    assert result.filas_aceptadas == 1
    assert result.rows[0].call_sign == "LPE2000"


def test_hora_como_texto_hhmm_se_normaliza():
    """La hoja del FMP guarda la hora como texto '00:05'."""
    content = _build_season_xlsx([
        ("LPE2000", "D", 174, "A320", "TYL", "TYL", date(2026, 9, 9), "00:05", "J"),
        ("JAP7221", "A", 186, "A320", "PIU", "PIU", date(2026, 9, 9), "23:50:00", "J"),
    ])
    result = parse_itinerary_season_excel(content, IATA_TO_ICAO)
    assert result.filas_rechazadas == 0
    assert [r.hora_utc for r in result.rows] == ["0005", "2350"]


def test_elige_la_hoja_que_cubre_la_fecha_y_no_la_cruda():
    """Del libro real: la hoja convertida quedo en junio, la cruda trae el
    numero de vuelo sin aerolinea, y la buena es la del FMP."""
    content = _libro({
        "S26 - formato CONVERTIDO": (HEADERS_TEMPORADA, [
            ("LPE2000", "D", 174, "A320", "TYL", "TYL", date(2026, 6, 30), "00:05", "J"),
        ]),
        "Formato DGAC para convertir": (HEADERS_TEMPORADA, [
            ("2255", "D", 174, "A320", "TYL", "TYL", date(2026, 9, 9), "00:05", "J"),
        ]),
        "formato segregado para FMP": (HEADERS_FMP, [
            ("D", "LPE2255", "A320", "SPYL", date(2026, 9, 9), "00:05", "J"),
        ]),
    })
    result = parse_itinerary_season_excel(content, IATA_TO_ICAO, min_date=date(2026, 9, 9))
    assert result.hoja == "formato segregado para FMP"
    assert [h.nombre for h in result.hojas].count("S26 - formato CONVERTIDO") == 1
    assert result.filas_aceptadas == 1
    assert result.rows[0].call_sign == "LPE2255"
    assert result.rows[0].direction == "DEP"


def test_hoja_elegida_a_mano_manda():
    content = _libro({
        "buena": (HEADERS_FMP, [("D", "LPE2255", "A320", "SPYL", date(2026, 9, 9), "00:05", "J")]),
        "otra": (HEADERS_FMP, [("A", "JAP7221", "A320", "SPZO", date(2026, 9, 9), "01:00", "J")]),
    })
    result = parse_itinerary_season_excel(
        content, IATA_TO_ICAO, min_date=date(2026, 9, 9), hoja="otra"
    )
    assert result.hoja == "otra"
    assert result.rows[0].call_sign == "JAP7221"

    with pytest.raises(InvalidItineraryFile, match="no esta en el archivo"):
        parse_itinerary_season_excel(content, IATA_TO_ICAO, hoja="inexistente")


def test_archivo_sin_columna_de_direccion_se_rechaza_con_motivo():
    """El recorte pegado sin la columna D/A: no hay forma de saber si cada
    vuelo es arribo o despegue, asi que se rechaza entero."""
    headers = [h for h in HEADERS_FMP if h != "D / A"]
    content = _libro({"Hoja1": (headers, [
        ("LPE2255", "A320", "SPYL", date(2026, 9, 9), "00:05", "J"),
    ])})
    with pytest.raises(InvalidItineraryFile, match="columna de dirección"):
        parse_itinerary_season_excel(content, IATA_TO_ICAO, min_date=date(2026, 9, 9))


def test_ninguna_hoja_llega_a_la_fecha_pedida():
    content = _libro({"Hoja1": (HEADERS_TEMPORADA, [
        ("LPE2000", "D", 174, "A320", "TYL", "TYL", date(2026, 6, 30), "00:05", "J"),
    ])})
    with pytest.raises(InvalidItineraryFile, match="Ninguna hoja"):
        parse_itinerary_season_excel(content, IATA_TO_ICAO, min_date=date(2026, 9, 9))


def _con_formulas(content: bytes, enlaces_externos: bool) -> bytes:
    """El mismo xlsx pero con las celdas convertidas en formulas, conservando
    el valor cacheado.

    Es lo que Excel guarda cuando el itinerario se arma enlazando otra
    planilla: en pantalla se ve lo recalculado, pero lo que queda escrito es
    el ultimo valor guardado, que puede ser de otra fecha. Con
    `enlaces_externos` se elige si esas formulas apuntan a otro libro (hay que
    rechazar el archivo) o son del mismo (normal: no se toca)."""
    import io
    import re
    import zipfile

    entrada = zipfile.ZipFile(io.BytesIO(content))
    buf = io.BytesIO()
    referencia = "'[1]Origen'!$A2" if enlaces_externos else "Hoja1!$A2"
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as salida:
        for item in entrada.infolist():
            datos = entrada.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = datos.decode("utf-8")
                datos = re.sub("<v>", "<f>" + referencia + "</f><v>", xml).encode("utf-8")
            salida.writestr(item, datos)
        if enlaces_externos:
            salida.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")
    return buf.getvalue()


def test_rechaza_el_archivo_que_es_solo_formulas_a_otro_libro():
    """Caso real (set-2026): un recorte "del 9 al 11 de setiembre" cuyas celdas
    eran formulas al libro de la DGAC. Excel mostraba setiembre, pero adentro
    tenia guardado el 1 al 4 de julio."""
    from datetime import time as time_

    base = _build_season_xlsx([
        ("LPE2000", "D", 174, "A320", "TYL", "TYL", date(2026, 1, 1), time_(0, 0), "J"),
    ])

    with pytest.raises(InvalidItineraryFile) as excepcion:
        parse_itinerary_season_excel(
            _con_formulas(base, enlaces_externos=True), IATA_TO_ICAO, min_date=date(2026, 1, 1)
        )
    assert "otro libro" in str(excepcion.value)

    # Las formulas internas son normales (el libro de la DGAC las usa) y no
    # deben rechazar nada: el valor guardado es del mismo archivo.
    resultado = parse_itinerary_season_excel(
        _con_formulas(base, enlaces_externos=False), IATA_TO_ICAO, min_date=date(2026, 1, 1)
    )
    assert resultado.filas_aceptadas == 1
